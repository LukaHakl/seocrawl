"""Keyword-to-URL matching: the join that turns a keyword list into a map.

A keyword list on its own is a wish. Joined against a crawl it becomes four
actionable buckets: which URL **owns** a term, which terms have **no owner**
(the money rows), where two pages **compete** for the same term, and where the
targeting is **inverted** -- a head term pointed at a single product, or a
long-tail term pointed at a category.
"""

from __future__ import annotations

import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence
from urllib.parse import urlsplit

from openpyxl import load_workbook

from .kwplanner import KeywordRow

# --- tokenizing -------------------------------------------------------------

#: Words that carry no targeting signal. Kept deliberately short: every word
#: removed here is a word that can no longer distinguish two pages.
STOPWORDS = frozenset(
    {
        "the", "for", "and", "a", "an", "of", "to", "in", "on", "with",
        "buy", "shop", "online", "sale",
        # URL furniture, not content
        "collections", "collection", "products", "product", "pages", "page",
        "blogs", "blog", "www", "html", "index",
    }
)

#: Field weights: what a page *declares* it is about outranks its body copy.
SLUG_WEIGHT = 3
TITLE_WEIGHT = 2
H1_WEIGHT = 2
MAX_FIELD_WEIGHT = max(SLUG_WEIGHT, TITLE_WEIGHT, H1_WEIGHT)

#: Classification thresholds.
OWN_THRESHOLD = 0.75
WEAK_THRESHOLD = 0.50
CONTENDER_MARGIN = 0.25

OWNED, WEAK, GAP, CANNIBAL = "OWNED", "WEAK", "GAP", "CANNIBAL"

#: A keyword of this many tokens or more is long-tail.
LONG_TAIL_TOKENS = 4
#: Head keywords are those at or above this percentile of the file's volume.
HEAD_PERCENTILE = 75

_WORD = re.compile(r"[a-z0-9]+")


def singularize(word: str) -> str:
    """Naive plural stripping: ``wallets`` -> ``wallet``.

    Deliberately naive. It only has to make "wallet" and "wallets" collide;
    getting "glasses" right is not worth the failure modes of a real stemmer.
    """
    if len(word) > 3 and word.endswith("s") and not word.endswith(("ss", "us", "is")):
        return word[:-1]
    return word


def tokenize_target(text: str, *, stopwords: Iterable[str] = STOPWORDS) -> list[str]:
    """Lowercase, split, drop stopwords, singularize."""
    stop = set(stopwords)
    return [
        singular
        for word in _WORD.findall((text or "").lower())
        if (singular := singularize(word)) not in stop and len(singular) > 1
    ]


def slug_tokens(url: str) -> list[str]:
    """Tokens from a URL's path, with the structural segments removed."""
    path = urlsplit(url).path
    return tokenize_target(re.sub(r"[/_\-.]+", " ", path))


def page_type(url: str) -> str:
    """Infer what kind of page a URL is from its Shopify-style path."""
    path = urlsplit(url).path.lower()
    if "/products/" in path:
        return "product"
    if "/collections/" in path:
        return "collection"
    if "/blogs/" in path or "/blog/" in path:
        return "content"
    if "/pages/" in path:
        return "static"
    return "home" if path in ("", "/") else "other"


# --- pages ------------------------------------------------------------------


@dataclass
class PageTarget:
    """One crawled page reduced to what it declares it is about."""

    url: str
    title: str = ""
    h1: str = ""
    page_type: str = "other"
    product_group_id: str | None = None
    inlinks: int = 0

    slug: tuple[str, ...] = field(default_factory=tuple)
    title_tokens: tuple[str, ...] = field(default_factory=tuple)
    h1_tokens: tuple[str, ...] = field(default_factory=tuple)

    def __post_init__(self) -> None:
        if not self.slug:
            self.slug = tuple(slug_tokens(self.url))
        if not self.title_tokens:
            self.title_tokens = tuple(tokenize_target(self.title))
        if not self.h1_tokens:
            self.h1_tokens = tuple(tokenize_target(self.h1))
        if self.page_type == "other":
            self.page_type = page_type(self.url)

    @property
    def all_tokens(self) -> set[str]:
        return set(self.slug) | set(self.title_tokens) | set(self.h1_tokens)

    def field_weight(self, token: str) -> int:
        """The strongest field this token appears in, as a weight."""
        if token in self.slug:
            return SLUG_WEIGHT
        if token in self.title_tokens or token in self.h1_tokens:
            return max(TITLE_WEIGHT, H1_WEIGHT)
        return 0


#: Pages sheet headers -> PageTarget fields. Matching by header text rather
#: than position keeps this working when columns are added to the sheet.
_SHEET_COLUMNS = {
    "url": "url",
    "title": "title",
    "h1": "h1",
    "product group": "product_group_id",
    "inlinks": "inlinks",
    "status": "_status",
    "indexable": "_indexable",
}


def load_pages_from_xlsx(path: Path | str, *, indexable_only: bool = True) -> list[PageTarget]:
    """Read the Pages sheet of an audit workbook into match targets.

    Only indexable 200s are loaded by default: a keyword cannot be owned by a
    page that redirects, 404s or carries noindex, and letting one win the match
    would hide the real gap behind a page that can never rank.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Pages" not in workbook.sheetnames:
            raise ValueError("%s has no 'Pages' sheet -- is it a seocrawl audit?" % path)
        sheet = workbook["Pages"]

        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []

        index: dict[str, int] = {}
        for position, cell in enumerate(header):
            key = _SHEET_COLUMNS.get(str(cell or "").strip().lower())
            if key and key not in index:
                index[key] = position

        if "url" not in index:
            raise ValueError("%s: Pages sheet has no URL column." % path)

        def value(row, key):
            position = index.get(key)
            if position is None or position >= len(row):
                return None
            return row[position]

        pages: list[PageTarget] = []
        for row in rows:
            url = value(row, "url")
            if not url:
                continue
            if indexable_only:
                status = value(row, "_status")
                if status is not None and status != 200:
                    continue
                if str(value(row, "_indexable") or "yes").lower() == "no":
                    continue

            inlinks = value(row, "inlinks")
            pages.append(
                PageTarget(
                    url=str(url),
                    title=str(value(row, "title") or ""),
                    h1=str(value(row, "h1") or ""),
                    product_group_id=(str(value(row, "product_group_id")).strip() or None)
                    if value(row, "product_group_id") else None,
                    inlinks=int(inlinks) if isinstance(inlinks, (int, float)) else 0,
                )
            )
        return pages
    finally:
        workbook.close()


def pages_from_crawl(results: Iterable) -> list[PageTarget]:
    """Build match targets straight from a live crawl's ``PageResult`` objects."""
    return [
        PageTarget(url=r.url, title=r.title or "", h1=r.h1 or "",
                   product_group_id=r.product_group_id, inlinks=r.inlinks)
        for r in results if r.indexable
    ]


# --- scoring ----------------------------------------------------------------


class TokenInformativeness:
    """How much each token distinguishes one page from another.

    On a single-niche site almost every page says "leather" and "wallet", so
    those tokens separate nothing: a keyword scored on raw token overlap would
    match half the catalogue and the gap list would come back empty. Weighting
    each token by inverse document frequency fixes that -- common tokens fade,
    the distinguishing token ("minimalist", "bifold") decides the match.

    With ``enabled=False`` every token weighs the same, which is the plain
    weighted-coverage behaviour.
    """

    #: No token ever drops to zero weight. Without a floor, a term appearing on
    #: every page counts for nothing, and a three-word keyword can end up
    #: decided by a single token -- which on a small crawl is close to a coin
    #: toss. The floor keeps common tokens contributing while still letting the
    #: distinguishing one dominate.
    FLOOR = 0.25

    def __init__(self, pages: Sequence[PageTarget], *, enabled: bool = True) -> None:
        self.enabled = enabled
        self.page_count = len(pages)
        self.document_frequency: Counter[str] = Counter()
        for page in pages:
            self.document_frequency.update(page.all_tokens)
        self._ceiling = math.log(self.page_count + 1) or 1.0

    def weight(self, token: str) -> float:
        """Informativeness of ``token`` in [FLOOR, 1]; 1.0 when disabled."""
        if not self.enabled or not self.page_count:
            return 1.0
        frequency = self.document_frequency.get(token, 0)
        rarity = math.log((self.page_count + 1) / (frequency + 1)) / self._ceiling
        return self.FLOOR + (1.0 - self.FLOOR) * max(0.0, min(1.0, rarity))


def score_page(
    tokens: Sequence[str], page: PageTarget, informativeness: TokenInformativeness
) -> float:
    """Weighted fraction of a keyword's tokens that the page targets.

    Each token contributes the weight of the strongest field it appears in
    (slug 3, title/H1 2), scaled by how informative the token is. A page whose
    slug contains every informative token scores 1.0.
    """
    if not tokens:
        return 0.0

    weights = [informativeness.weight(token) for token in tokens]
    total = sum(weights)
    if total <= 0:
        # Every token is generic; fall back to plain coverage so the keyword is
        # still judged on something rather than dividing by zero.
        weights = [1.0] * len(tokens)
        total = float(len(tokens))

    earned = sum(
        weight * page.field_weight(token)
        for token, weight in zip(tokens, weights)
    )
    return earned / (total * MAX_FIELD_WEIGHT)


# --- results ----------------------------------------------------------------


@dataclass
class Contender:
    """One page competing for a keyword."""

    url: str
    page_type: str
    score: float


@dataclass
class KeywordMatch:
    """One keyword's verdict against the crawl."""

    row: KeywordRow
    status: str
    score: float = 0.0
    owner: str | None = None
    owner_page_type: str | None = None
    contenders: list[Contender] = field(default_factory=list)
    inverted: str = ""

    @property
    def keyword(self) -> str:
        return self.row.keyword

    @property
    def volume_mid(self) -> int:
        return self.row.volume_mid

    @property
    def contender_summary(self) -> str:
        return " | ".join("%s (%.2f)" % (c.url, c.score) for c in self.contenders)


@dataclass
class GapCluster:
    """A group of unowned keywords sharing a head phrase."""

    head: str
    keywords: list[KeywordMatch] = field(default_factory=list)

    @property
    def volume(self) -> int:
        return sum(k.volume_mid for k in self.keywords)

    @property
    def size(self) -> int:
        return len(self.keywords)

    def suggested_action(self) -> str:
        """The one-line recommendation printed in the Gaps sheet."""
        return (
            "%d keyword%s, combined volume_mid %s - no owning URL; "
            "candidate: new collection or landing page targeting '%s'"
            % (self.size, "" if self.size == 1 else "s", "{:,}".format(self.volume),
               self.head)
        )


@dataclass
class KeywordMap:
    """Everything the join produced."""

    matches: list[KeywordMatch] = field(default_factory=list)
    clusters: list[GapCluster] = field(default_factory=list)
    page_count: int = 0
    head_volume_threshold: int = 0
    idf_enabled: bool = True

    def by_status(self, status: str) -> list[KeywordMatch]:
        return [m for m in self.matches if m.status == status]

    @property
    def counts(self) -> dict[str, int]:
        counter = Counter(m.status for m in self.matches)
        return {s: counter.get(s, 0) for s in (OWNED, WEAK, GAP, CANNIBAL)}


# --- the join ---------------------------------------------------------------


def _percentile(values: Sequence[int], percentile: float) -> int:
    """Nearest-rank percentile of a list of volumes."""
    if not values:
        return 0
    ordered = sorted(values)
    rank = max(0, min(len(ordered) - 1,
                      int(math.ceil(percentile / 100 * len(ordered))) - 1))
    return ordered[rank]


def _same_product_group(contenders: Sequence[Contender],
                        pages: dict[str, PageTarget]) -> bool:
    """True when every contender is a variant of one and the same product.

    Reuses the v1.1 ProductGroup identity: three URLs for one wallet in three
    colours are not competing with each other, and reporting them as
    cannibalization would bury the cases that are real.
    """
    groups = set()
    for contender in contenders:
        page = pages.get(contender.url)
        if page is None or not page.product_group_id:
            return False            # at least one contender is a different thing
        groups.add(page.product_group_id)
    return len(groups) == 1


def classify(
    tokens: Sequence[str],
    scored: Sequence[Contender],
    pages: dict[str, PageTarget],
) -> tuple[str, list[Contender]]:
    """Turn a keyword's page scores into a status plus its contenders."""
    if not scored:
        return GAP, []

    best = scored[0].score
    contenders = [c for c in scored
                  if c.score >= OWN_THRESHOLD and best - c.score <= CONTENDER_MARGIN]

    if len(contenders) >= 2 and not _same_product_group(contenders, pages):
        return CANNIBAL, contenders
    if best >= OWN_THRESHOLD:
        return OWNED, contenders[:1] or [scored[0]]
    if best >= WEAK_THRESHOLD:
        return WEAK, [scored[0]]
    return GAP, []


def detect_inversion(match: KeywordMatch, head_threshold: int) -> str:
    """Flag targeting that points the wrong kind of keyword at the wrong page.

    The healthy pattern is heads -> collections, specifics -> products,
    questions -> content. A head term owned by one product cannot rank for the
    breadth it implies; a five-word specific owned by a collection sends the
    shopper to a grid instead of the thing they asked for.
    """
    if match.status not in (OWNED, CANNIBAL) or not match.owner_page_type:
        return ""

    token_count = len(match.row.tokens)
    if match.volume_mid >= head_threshold and match.owner_page_type == "product":
        return "head keyword owned by a product page"
    if token_count >= LONG_TAIL_TOKENS and match.owner_page_type == "collection":
        return "long-tail keyword owned by a collection page"
    return ""


def map_keywords(
    rows: Sequence[KeywordRow],
    pages: Sequence[PageTarget],
    *,
    use_idf: bool = True,
) -> KeywordMap:
    """Join keywords against crawled pages and classify every one."""
    by_url = {page.url: page for page in pages}
    informativeness = TokenInformativeness(pages, enabled=use_idf)
    head_threshold = _percentile([r.volume_mid for r in rows], HEAD_PERCENTILE)

    matches: list[KeywordMatch] = []
    for row in rows:
        tokens = [singularize(t) for t in tokenize_target(row.normalized)]
        scored = sorted(
            (
                Contender(page.url, page.page_type, score_page(tokens, page, informativeness))
                for page in pages
            ),
            key=lambda c: (-c.score, c.url),
        )
        scored = [c for c in scored if c.score > 0]

        status, contenders = classify(tokens, scored, by_url)
        # For a cannibalized keyword the "owner" is the best-scoring contender;
        # the whole point of the row is that the ownership is disputed, so the
        # rest are listed alongside it.
        best = contenders[0] if contenders else None
        match = KeywordMatch(
            row=row,
            status=status,
            score=round(scored[0].score, 3) if scored else 0.0,
            owner=best.url if best else None,
            owner_page_type=best.page_type if best else None,
            contenders=contenders if status == CANNIBAL else [],
        )
        match.inverted = detect_inversion(match, head_threshold)
        matches.append(match)

    matches.sort(key=lambda m: (
        {GAP: 0, CANNIBAL: 1, WEAK: 2, OWNED: 3}[m.status], -m.volume_mid, m.keyword
    ))

    return KeywordMap(
        matches=matches,
        clusters=cluster_gaps([m for m in matches if m.status == GAP]),
        page_count=len(pages),
        head_volume_threshold=head_threshold,
        idf_enabled=use_idf,
    )


# --- gap clustering ---------------------------------------------------------


def cluster_gaps(gaps: Sequence[KeywordMatch], *, min_cluster: int = 2) -> list[GapCluster]:
    """Group unowned keywords by a shared two-token phrase.

    Deliberately dumb: take the most common adjacent token pair across all gap
    keywords, pull every keyword containing it into that cluster, repeat. A
    cluster head you can read off the sheet and argue with beats a clever
    embedding you have to trust.
    """
    remaining = list(gaps)
    clusters: list[GapCluster] = []

    def bigrams(match: KeywordMatch) -> list[str]:
        tokens = tokenize_target(match.row.normalized)
        return [" ".join(tokens[i:i + 2]) for i in range(len(tokens) - 1)]

    while remaining:
        counts: Counter[str] = Counter()
        for match in remaining:
            counts.update(set(bigrams(match)))

        if not counts:
            break
        head, frequency = counts.most_common(1)[0]
        if frequency < min_cluster:
            break

        members = [m for m in remaining if head in bigrams(m)]
        clusters.append(GapCluster(head=head, keywords=members))
        remaining = [m for m in remaining if m not in members]

    # Whatever is left has nothing in common with anything: one cluster each.
    clusters.extend(GapCluster(head=m.row.normalized, keywords=[m]) for m in remaining)
    clusters.sort(key=lambda c: (-c.volume, -c.size, c.head))
    return clusters


# --- the caveat -------------------------------------------------------------

VOLUME_CAVEAT = (
    "GKP volumes are ranges for non-spending accounts and geo-dependent - treat "
    "as relative sizes, not truth. GSC data supersedes this file the day you get "
    "access."
)
