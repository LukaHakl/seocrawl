"""Tests for the xlsx workbook, the HTML report and the crawl diff."""

from __future__ import annotations

import datetime as dt

import pytest
from openpyxl import Workbook, load_workbook

from seocrawl.crawler import OutLink, PageResult
from seocrawl.export import (
    FILL_HIGH,
    FILL_MED,
    PAGE_COLUMNS,
    diff_crawls,
    read_previous,
    report_name,
    write_diff_sheet,
    write_html_report,
    write_workbook,
)
from seocrawl.findings import HIGH, LOW, MED, Audit, Finding, anchor_report
from seocrawl.psi import PsiResult

SITE = "https://shop.test"


def page(path: str, **kwargs) -> PageResult:
    defaults = dict(
        url=SITE + path, status=200, final_url=SITE + path,
        canonical=SITE + path, canonical_state="self",
        title="Title %s" % path,
        meta_description="Description for %s, long enough to pass the check." % path,
        h1_count=1, h1="H1 %s" % path, word_count=600,
        content_hash="h%s" % path, simhash=abs(hash(path)) & ((1 << 64) - 1),
        viewport="width=device-width", inlinks=4, size_bytes=51200,
    )
    defaults.update(kwargs)
    return PageResult(**defaults)


@pytest.fixture
def crawl():
    """A small crawl with one of each interesting problem."""
    pages = [
        page("/"),
        page("/products/bifold", is_product=True, price_consistent=False,
             schema_types=["Product"], inlinks=12,
             prices={"og:price:amount": "179", "jsonld:offers.price": "169"}),
        page("/products/cardholder", is_product=True, schema_types=["Product"],
             review_suspicious=True, review_schema_count=4187, review_onpage_count=3,
             review_reason="schema claims 4187 reviews, page renders ~3"),
        page("/gone", status=404, title=None, title_issue="missing"),
        page("/old", redirect_hops=2, redirect_chain=["301 -> %s/mid" % SITE,
                                                      "301 -> %s/new" % SITE]),
    ]
    links = [
        OutLink(SITE + "/", SITE + "/gone", "Sale", True, False),
        OutLink(SITE + "/", SITE + "/products/bifold", "Bifold wallet", True, False),
        OutLink(SITE + "/products/cardholder", SITE + "/products/bifold",
                "Bifold wallet", True, False),
    ]
    audit = Audit({p.url: p for p in pages}, links, {SITE + "/", SITE + "/gone"})
    findings = audit.run()
    return pages, links, findings, audit.stats(findings)


# --- naming -----------------------------------------------------------------


def test_report_name_uses_domain_and_date():
    name = report_name("https://www.acmeleather.com/", "xlsx", dt.date(2026, 8, 26))
    assert name == "audit_acmeleather-com_2026-08-26.xlsx"


def test_report_name_handles_html():
    assert report_name("https://x.io", "html", dt.date(2026, 1, 2)) == "audit_x-io_2026-01-02.html"


# --- workbook ---------------------------------------------------------------


def test_workbook_has_all_sheets(crawl, tmp_path):
    pages, links, findings, stats = crawl
    path = write_workbook(tmp_path / "a.xlsx", pages, findings, stats,
                          anchor_report(links), [PsiResult(url=SITE + "/", lcp_ms=3000.0)],
                          site="shop.test")
    book = load_workbook(path)
    assert book.sheetnames == ["Summary", "Pages", "Findings", "Anchors", "PSI", "_data"]
    assert book["_data"].sheet_state == "hidden"


def test_optional_sheets_are_omitted_when_empty(crawl, tmp_path):
    pages, _, findings, stats = crawl
    book = load_workbook(write_workbook(tmp_path / "b.xlsx", pages, findings, stats))
    assert "Anchors" not in book.sheetnames and "PSI" not in book.sheetnames


def test_pages_sheet_has_one_row_per_url_and_a_frozen_header(crawl, tmp_path):
    pages, _, findings, stats = crawl
    sheet = load_workbook(write_workbook(tmp_path / "c.xlsx", pages, findings, stats))["Pages"]
    assert sheet.max_row == len(pages) + 1
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref.startswith("A1:")
    assert [c.value for c in sheet[1]] == [c[0] for c in PAGE_COLUMNS]


def test_pages_sheet_paints_high_issue_cells_red(crawl, tmp_path):
    """A 404 status and a price mismatch must both show up as red cells."""
    pages, _, findings, stats = crawl
    sheet = load_workbook(write_workbook(tmp_path / "d.xlsx", pages, findings, stats))["Pages"]

    headers = [c.value for c in sheet[1]]
    status_col = headers.index("Status") + 1
    price_col = headers.index("Price consistent") + 1

    reds = {sheet.cell(row=r, column=status_col).value
            for r in range(2, sheet.max_row + 1)
            if sheet.cell(row=r, column=status_col).fill.fgColor.rgb == FILL_HIGH.fgColor.rgb}
    assert 404 in reds

    price_reds = [r for r in range(2, sheet.max_row + 1)
                  if sheet.cell(row=r, column=price_col).fill.fgColor.rgb
                  == FILL_HIGH.fgColor.rgb]
    assert len(price_reds) == 1


def test_pages_sheet_paints_med_issues_amber(crawl, tmp_path):
    pages, _, findings, stats = crawl
    sheet = load_workbook(write_workbook(tmp_path / "e.xlsx", pages, findings, stats))["Pages"]
    headers = [c.value for c in sheet[1]]
    hops_col = headers.index("Redirect hops") + 1
    ambers = [r for r in range(2, sheet.max_row + 1)
              if sheet.cell(row=r, column=hops_col).fill.fgColor.rgb == FILL_MED.fgColor.rgb]
    assert len(ambers) == 1


def test_findings_sheet_is_severity_sorted(crawl, tmp_path):
    pages, _, findings, stats = crawl
    sheet = load_workbook(write_workbook(tmp_path / "f.xlsx", pages, findings, stats))["Findings"]
    severities = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    order = {HIGH: 0, MED: 1, LOW: 2}
    assert severities == sorted(severities, key=lambda s: order[s])


def test_findings_sheet_carries_why_it_matters(crawl, tmp_path):
    pages, _, findings, stats = crawl
    sheet = load_workbook(write_workbook(tmp_path / "g.xlsx", pages, findings, stats))["Findings"]
    whys = [sheet.cell(row=r, column=6).value for r in range(2, sheet.max_row + 1)]
    assert all(w and len(w) > 20 for w in whys)


def test_summary_sheet_names_the_site_and_top_findings(crawl, tmp_path):
    pages, _, findings, stats = crawl
    sheet = load_workbook(write_workbook(tmp_path / "h.xlsx", pages, findings, stats,
                                         site="shop.test"))["Summary"]
    text = "\n".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert "shop.test" in text
    assert "URLs crawled" in text and "Top findings" in text
    assert "Product price signals disagree" in text


def test_anchors_sheet_orders_by_usage(crawl, tmp_path):
    pages, links, findings, stats = crawl
    sheet = load_workbook(write_workbook(tmp_path / "i.xlsx", pages, findings, stats,
                                         anchor_report(links)))["Anchors"]
    assert sheet.cell(row=2, column=1).value == "Bifold wallet"
    assert sheet.cell(row=2, column=3).value == 2


def test_psi_sheet_highlights_failing_metrics(crawl, tmp_path):
    pages, _, findings, stats = crawl
    results = [PsiResult(url=SITE + "/", performance_score=0.35, lcp_ms=4800.0,
                         cls=0.3, tbt_ms=900.0, field_overall="SLOW")]
    sheet = load_workbook(write_workbook(tmp_path / "j.xlsx", pages, findings, stats,
                                         psi_results=results))["PSI"]
    headers = [c.value for c in sheet[1]]
    lcp = sheet.cell(row=2, column=headers.index("LCP (ms)") + 1)
    assert lcp.value == 4800.0 and lcp.fill.fgColor.rgb == FILL_HIGH.fgColor.rgb


def test_workbook_handles_an_empty_crawl(tmp_path):
    from seocrawl.findings import SiteStats
    path = write_workbook(tmp_path / "empty.xlsx", [], [], SiteStats())
    assert load_workbook(path)["Pages"].max_row == 1


# --- diff -------------------------------------------------------------------


def _write(tmp_path, name, pages, findings):
    from seocrawl.findings import SiteStats
    return write_workbook(tmp_path / name, pages, findings, SiteStats(total_urls=len(pages)))


def test_read_previous_round_trips_via_the_data_sheet(crawl, tmp_path):
    pages, _, findings, _ = crawl
    path = _write(tmp_path, "prev.xlsx", pages, findings)
    old_findings, old_urls = read_previous(path)
    assert {f.key for f in old_findings} == {f.key for f in findings}
    assert old_urls == {p.url for p in pages}


def test_read_previous_falls_back_to_visible_sheets(crawl, tmp_path):
    """A workbook without the hidden _data sheet must still be diffable."""
    pages, _, findings, _ = crawl
    path = _write(tmp_path, "old_format.xlsx", pages, findings)

    book = load_workbook(path)
    del book["_data"]
    book.save(path)

    old_findings, old_urls = read_previous(path)
    assert old_findings and old_urls == {p.url for p in pages}


def test_diff_detects_new_and_fixed_issues(crawl, tmp_path):
    pages, _, findings, _ = crawl
    previous = _write(tmp_path, "before.xlsx", pages, findings)

    # This crawl: the price mismatch is fixed, a new server error appeared.
    kept = [f for f in findings if "price signals" not in f.issue]
    new = kept + [Finding(HIGH, "indexability", "Server errors (5xx)",
                          [SITE + "/boom"], "1 URL", "why")]
    diff = diff_crawls(previous, new, pages)

    assert any("Server errors" in f.issue for f in diff.new_issues)
    assert any("price signals" in f.issue for f in diff.fixed_issues)


def test_diff_detects_counts_getting_worse(crawl, tmp_path):
    pages, _, findings, _ = crawl
    previous = _write(tmp_path, "before2.xlsx", pages, findings)

    grown = [Finding(f.severity, f.category, f.issue,
                     f.urls + [SITE + "/extra%d" % i for i in range(3)], f.detail, f.why_it_matters)
             for f in findings]
    diff = diff_crawls(previous, grown, pages)
    assert diff.worsened and not diff.improved
    finding, was, now = diff.worsened[0]
    assert now == was + 3


def test_diff_detects_new_and_removed_urls(crawl, tmp_path):
    pages, _, findings, _ = crawl
    previous = _write(tmp_path, "before3.xlsx", pages, findings)

    current = [p for p in pages if not p.url.endswith("/gone")] + [page("/new-page")]
    diff = diff_crawls(previous, findings, current)
    assert diff.new_urls == [SITE + "/new-page"]
    assert diff.removed_urls == [SITE + "/gone"]


def test_identical_crawls_diff_to_nothing(crawl, tmp_path):
    pages, _, findings, _ = crawl
    previous = _write(tmp_path, "same.xlsx", pages, findings)
    assert diff_crawls(previous, findings, pages).is_empty


def test_diff_sheet_is_written(crawl, tmp_path):
    pages, _, findings, _ = crawl
    previous = _write(tmp_path, "before4.xlsx", pages, findings)
    diff = diff_crawls(previous, [f for f in findings if "price" not in f.issue], pages)

    book = Workbook()
    write_diff_sheet(book, diff)
    text = "\n".join(str(c.value) for row in book["Diff"].iter_rows() for c in row if c.value)
    assert "FIXED" in text and book.sheetnames[1] == "Diff"


# --- HTML -------------------------------------------------------------------


def test_html_report_is_self_contained(crawl, tmp_path):
    pages, _, findings, stats = crawl
    path = write_html_report(tmp_path / "r.html", stats, findings, site="shop.test")
    document = path.read_text(encoding="utf-8")

    assert document.startswith("<!doctype html>")
    assert "<style>" in document
    assert "src=\"http" not in document and "<script" not in document


def test_html_report_shows_findings_and_why(crawl, tmp_path):
    pages, _, findings, stats = crawl
    document = write_html_report(tmp_path / "r2.html", stats, findings,
                                 site="shop.test").read_text(encoding="utf-8")
    assert "Product price signals disagree" in document
    assert "Merchant Center" in document
    assert "Why it matters" in document


def test_html_report_escapes_page_content(tmp_path):
    from seocrawl.findings import SiteStats
    finding = Finding(HIGH, "on-page", "Bad <script>alert(1)</script> title",
                      ["https://x.test/<img>"], "detail & more", "why & how")
    document = write_html_report(tmp_path / "r3.html", SiteStats(), [finding],
                                 site="x.test").read_text(encoding="utf-8")
    assert "<script>alert(1)</script>" not in document
    assert "&lt;script&gt;" in document


def test_html_report_includes_psi_and_diff(crawl, tmp_path):
    pages, _, findings, stats = crawl
    previous = _write(tmp_path, "before5.xlsx", pages, findings)
    diff = diff_crawls(previous, [f for f in findings if "price" not in f.issue], pages)
    document = write_html_report(
        tmp_path / "r4.html", stats, findings, site="shop.test",
        psi_results=[PsiResult(url=SITE + "/", performance_score=0.4, lcp_ms=4000.0)],
        diff=diff,
    ).read_text(encoding="utf-8")
    assert "PageSpeed (mobile)" in document
    assert "Change since last crawl" in document
    assert "4.0s" in document


def test_html_report_handles_a_clean_site(tmp_path):
    from seocrawl.findings import SiteStats
    document = write_html_report(tmp_path / "clean.html", SiteStats(total_urls=5), [],
                                 site="clean.test").read_text(encoding="utf-8")
    assert "No issues found" in document
