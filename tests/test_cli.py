"""Tests for config resolution and the CLI, including one end-to-end run."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from seocrawl.cli import build_parser, main, resolve_config
from seocrawl.config import (DEFAULT_CONTACT, DEFAULT_USER_AGENT,
                             PLACEHOLDER_CONTACT, Config)

# Reuse the fake site from the crawler tests for the end-to-end run.
from .test_crawler import Handler, site  # noqa: F401


@pytest.fixture(autouse=True)
def isolated_cwd(tmp_path, monkeypatch):
    """Run every test from an empty directory.

    ``Config.load`` falls back to ``./seocrawl.toml``, so without this the repo's
    own config file would silently feed into tests that mean to exercise the
    defaults.
    """
    workdir = tmp_path / "cwd"
    workdir.mkdir()
    monkeypatch.chdir(workdir)


# --- config -----------------------------------------------------------------


def test_defaults_are_polite():
    config = Config(start_url="https://x.test")
    assert config.delay == 0.7 and config.workers == 4
    assert config.respect_robots and config.max_urls == 10_000


def test_user_agent_carries_a_contact_address():
    assert "mailto:" in DEFAULT_USER_AGENT
    assert PLACEHOLDER_CONTACT in DEFAULT_USER_AGENT
    assert "seocrawl" in DEFAULT_USER_AGENT


def test_default_contact_is_a_placeholder_not_a_real_address():
    """Shipping a real address would point complaints at whoever wrote this
    rather than whoever ran it, so the default must stay unusable."""
    assert DEFAULT_CONTACT == PLACEHOLDER_CONTACT
    assert DEFAULT_CONTACT.endswith("@example.com")


def test_include_exclude_patterns():
    config = Config(start_url="https://x.test", include=[r"/products/"],
                    exclude=[r"\?variant="])
    assert config.allows("https://x.test/products/a")
    assert not config.allows("https://x.test/pages/about")
    assert not config.allows("https://x.test/products/a?variant=1")


def test_empty_include_allows_everything():
    assert Config(start_url="https://x.test").allows("https://x.test/anything")


def test_toml_is_loaded(tmp_path):
    path = tmp_path / "seocrawl.toml"
    path.write_text(
        '[seocrawl]\nmax_urls = 42\ndelay = 2.5\nexclude = ["/cart"]\n', encoding="utf-8"
    )
    config = Config.load(path, start_url="https://x.test")
    assert config.max_urls == 42 and config.delay == 2.5
    assert not config.allows("https://x.test/cart")


def test_flat_toml_without_a_section_works(tmp_path):
    path = tmp_path / "flat.toml"
    path.write_text("max_urls = 7\n", encoding="utf-8")
    assert Config.load(path, start_url="https://x.test").max_urls == 7


def test_toml_ignores_unknown_keys(tmp_path):
    path = tmp_path / "extra.toml"
    path.write_text("[seocrawl]\nmax_urls = 5\nnot_a_real_key = true\n", encoding="utf-8")
    assert Config.load(path, start_url="https://x.test").max_urls == 5


def test_cli_flag_beats_the_toml_file(tmp_path):
    path = tmp_path / "seocrawl.toml"
    path.write_text("[seocrawl]\nmax_urls = 42\ndelay = 2.5\n", encoding="utf-8")

    args = build_parser().parse_args(["https://x.test", "--config", str(path),
                                      "--max-urls", "9"])
    config = resolve_config(args)
    assert config.max_urls == 9      # flag wins
    assert config.delay == 2.5       # file still supplies the rest


def test_unset_flags_do_not_clobber_the_file(tmp_path):
    path = tmp_path / "seocrawl.toml"
    path.write_text("[seocrawl]\nworkers = 8\n", encoding="utf-8")
    args = build_parser().parse_args(["https://x.test", "--config", str(path)])
    assert resolve_config(args).workers == 8


def test_start_url_is_normalized():
    args = build_parser().parse_args(["HTTP://Example.com/Path/#frag"])
    assert resolve_config(args).start_url == "http://example.com/Path"


def test_psi_key_falls_back_to_the_environment(monkeypatch):
    monkeypatch.setenv("PSI_API_KEY", "from-env")
    args = build_parser().parse_args(["https://x.test", "--psi"])
    assert resolve_config(args).psi_api_key == "from-env"


def test_explicit_psi_key_beats_the_environment(monkeypatch):
    monkeypatch.setenv("PSI_API_KEY", "from-env")
    args = build_parser().parse_args(["https://x.test", "--psi", "--psi-key", "explicit"])
    assert resolve_config(args).psi_api_key == "explicit"


def test_boolean_off_switches():
    args = build_parser().parse_args(
        ["https://x.test", "--no-sitemap", "--ignore-robots", "--sitemap-only"]
    )
    config = resolve_config(args)
    assert not config.use_sitemap and not config.respect_robots and not config.follow_links


def test_repeatable_pattern_flags():
    args = build_parser().parse_args(
        ["https://x.test", "--exclude", "/a", "--exclude", "/b", "--include", "/p"]
    )
    config = resolve_config(args)
    assert config.exclude == ["/a", "/b"] and config.include == ["/p"]


# --- end to end -------------------------------------------------------------


def test_no_url_prints_help_and_exits_2(capsys):
    assert main([]) == 2
    assert "usage: seocrawl" in capsys.readouterr().out


def test_full_run_writes_both_reports(site, tmp_path, capsys):  # noqa: F811
    exit_code = main([
        site,
        "--delay", "0",
        "--workers", "2",
        "--output-dir", str(tmp_path),
        "--cache-dir", str(tmp_path / "cache"),
        "--recrawl",
        "--html",
    ])
    assert exit_code == 0

    xlsx = list(tmp_path.glob("audit_*.xlsx"))
    html = list(tmp_path.glob("audit_*.html"))
    assert len(xlsx) == 1 and len(html) == 1

    book = load_workbook(xlsx[0])
    assert {"Summary", "Pages", "Findings"} <= set(book.sheetnames)
    assert book["Pages"].max_row > 1

    output = capsys.readouterr().out
    assert "Crawl summary" in output and "Report:" in output


def test_run_then_diff_against_it(site, tmp_path):  # noqa: F811
    common = ["--delay", "0", "--workers", "2", "--cache-dir", str(tmp_path / "c")]
    assert main([site, *common, "--output-dir", str(tmp_path / "first"), "--recrawl"]) == 0
    previous = next((tmp_path / "first").glob("audit_*.xlsx"))

    assert main([site, *common, "--output-dir", str(tmp_path / "second"),
                 "--recrawl", "--diff", str(previous)]) == 0

    book = load_workbook(next((tmp_path / "second").glob("audit_*.xlsx")))
    assert "Diff" in book.sheetnames


def test_quiet_mode_still_writes_the_report(site, tmp_path):  # noqa: F811
    assert main([site, "--delay", "0", "--quiet", "--recrawl",
                 "--output-dir", str(tmp_path), "--cache-dir", str(tmp_path / "c")]) == 0
    assert list(tmp_path.glob("audit_*.xlsx"))


def test_missing_diff_file_is_a_warning_not_a_crash(site, tmp_path, capsys):  # noqa: F811
    assert main([site, "--delay", "0", "--recrawl", "--output-dir", str(tmp_path),
                 "--cache-dir", str(tmp_path / "c"),
                 "--diff", str(tmp_path / "nope.xlsx")]) == 0
    assert "not found" in capsys.readouterr().out


def test_malformed_config_reports_cleanly(tmp_path, capsys):
    """A broken TOML file must not surface as a tomllib traceback."""
    path = tmp_path / "broken.toml"
    path.write_text('[seocrawl]\nexclude = ["/search\\?"]\n', encoding="utf-8")
    assert main(["https://x.test", "--config", str(path)]) == 2
    assert "Config error" in capsys.readouterr().out


def test_shipped_config_is_valid_toml():
    """The seocrawl.toml shipped with the project must actually parse."""
    shipped = Path(__file__).resolve().parent.parent / "seocrawl.toml"
    values = Config.from_toml(shipped)
    assert values["max_urls"] == 10000
    assert Config.unknown_keys(shipped) == []

    config = Config(start_url="https://shop.test", **values)
    assert not config.allows("https://shop.test/cart")
    assert not config.allows("https://shop.test/c/all?filter.v.price=10")
    assert config.allows("https://shop.test/products/bifold")
