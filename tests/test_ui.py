"""Tests for the local web UI: routing, job lifecycle and the crawl it drives."""

from __future__ import annotations

import json
import threading
import time
from http.server import HTTPServer
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import pytest

from seocrawl.ui import Handler, Job, Runner, crawl_job
from seocrawl.ui_page import PAGE

# Reuse the fake site the crawler tests already serve.
from .test_crawler import Handler as SiteHandler


@pytest.fixture(scope="module")
def site():
    server = HTTPServer(("127.0.0.1", 0), SiteHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield "http://127.0.0.1:%d/" % server.server_port
    server.shutdown()


def wait_for_idle(runner, timeout=90):
    deadline = time.time() + timeout
    while runner.job.state == "running" and time.time() < deadline:
        time.sleep(0.2)
    return runner.job


@pytest.fixture
def ui(tmp_path):
    """The UI server, on its own port, writing into a temp directory.

    Teardown waits for any job still running: crawls live in daemon threads, so
    without this one test's crawl keeps writing while the next test's fixture has
    already swapped the runner out from under it.
    """
    runner = Runner(tmp_path)
    Handler.runner = runner
    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield "http://127.0.0.1:%d" % server.server_port, runner, tmp_path
    finally:
        runner.stop()
        wait_for_idle(runner, timeout=30)
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


#: Windows aborts a loopback connection with WinError 10053 often enough to
#: matter when the whole suite runs at once: several ThreadingHTTPServers are
#: alive, a crawl thread is saturating another, and the socket gets reset while
#: we are reading the response. It is a harness race, not a server bug.
#:
#: Only GETs are retried. Retrying a POST would be wrong here -- /api/stop is
#: not idempotent, so a retry after a *successful* first call returns
#: {"stopped": false} and looks identical to a genuine failure. Callers that
#: POST and care about the reset handle it themselves.
TRANSIENT = (ConnectionAbortedError, ConnectionResetError)


def get(base, path):
    for attempt in range(2):
        try:
            with urlopen(base + path, timeout=10) as response:
                body = response.read()
                return json.loads(body) if path.startswith("/api") else body
        except TRANSIENT:
            if attempt:
                raise
            time.sleep(0.2)


def post(base, path, payload=None, raw=None):
    body = raw if raw is not None else json.dumps(payload or {}).encode()
    try:
        with urlopen(Request(base + path, data=body, method="POST"), timeout=30) as r:
            return json.loads(r.read())
    except HTTPError as exc:
        return json.loads(exc.read())


# --- the page ---------------------------------------------------------------


def test_page_is_self_contained():
    """No CDN, no build step: the UI must work with no network at all."""
    assert PAGE.startswith("<!doctype html>")
    assert "<style>" in PAGE and "<script>" in PAGE
    assert 'src="http' not in PAGE and "cdn." not in PAGE


def test_page_is_served_at_root(ui):
    base, _runner, _dir = ui
    body = get(base, "/")
    assert b"seocrawl" in body and b"<!doctype html>" in body


def test_unknown_route_404s(ui):
    base, _runner, _dir = ui
    with pytest.raises(HTTPError) as excinfo:
        urlopen(base + "/nope", timeout=10)
    assert excinfo.value.code == 404


def test_page_declares_the_volume_caveat():
    """The caveat has to be visible wherever keyword volumes are, UI included."""
    assert "relative sizes, not truth" in PAGE
    assert "GSC data supersedes" in PAGE


# --- status -----------------------------------------------------------------


def test_status_starts_idle(ui):
    base, _runner, _dir = ui
    job = get(base, "/api/status")
    assert job["state"] == "idle" and job["done"] == 0


def test_job_serializes_every_field_the_page_reads():
    data = Job().as_dict()
    for key in ("state", "done", "total", "current", "message", "error",
                "elapsed", "outputs", "stats", "findings", "funnel"):
        assert key in data


# --- crawling ---------------------------------------------------------------


def test_crawl_runs_and_writes_both_reports(ui, site):
    base, runner, directory = ui
    started = post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 2,
                                        "spellcheck": False})
    assert started["started"] is True

    job = wait_for_idle(runner)
    assert job.state == "done", job.error
    assert list(directory.glob("audit_*.xlsx")) and list(directory.glob("audit_*.html"))
    assert job.stats["URLs crawled"] > 0
    assert len(job.outputs) == 2


def test_crawl_reports_progress(ui, site):
    base, runner, _dir = ui
    seen = []

    post(base, "/api/crawl", {"url": site, "delay": 0.05, "workers": 1,
                              "spellcheck": False})
    deadline = time.time() + 60
    while runner.job.state == "running" and time.time() < deadline:
        seen.append(get(base, "/api/status")["done"])
        time.sleep(0.1)
    wait_for_idle(runner)

    assert max(seen or [0]) > 0
    assert seen == sorted(seen)          # progress only ever moves forward


def test_findings_reach_the_page(ui, site):
    base, runner, _dir = ui
    post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 2,
                              "spellcheck": False})
    job = wait_for_idle(runner)
    assert job.findings
    for finding in job.findings:
        assert finding["severity"] in ("HIGH", "MED", "LOW")
        assert finding["issue"] and finding["why"]


def test_only_one_job_runs_at_a_time(ui, site):
    """Two concurrent crawls would share nothing but the target's patience."""
    base, runner, _dir = ui
    post(base, "/api/crawl", {"url": site, "delay": 0.2, "workers": 1,
                              "spellcheck": False})
    second = post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 1})
    assert second["started"] is False and "already running" in second["error"]
    runner.stop()
    wait_for_idle(runner)


def test_missing_url_is_reported_not_crashed(ui):
    base, runner, _dir = ui
    post(base, "/api/crawl", {"url": "  "})
    job = wait_for_idle(runner)
    assert job.state == "error" and "Enter a site URL" in job.error


def test_stop_ends_the_crawl_and_still_writes_a_report(ui, site):
    base, runner, directory = ui
    post(base, "/api/crawl", {"url": site, "delay": 0.3, "workers": 1,
                              "spellcheck": False})

    deadline = time.time() + 20
    while runner.job.done < 1 and time.time() < deadline:
        time.sleep(0.1)

    try:
        assert post(base, "/api/stop")["stopped"] is True
    except TRANSIENT:
        # The reset can land after the handler already stopped the crawl, so
        # losing the response says nothing about whether the stop took. Do not
        # retry -- a second /api/stop returns false whether the first worked or
        # not. The assertions below test the behaviour that actually matters.
        pass

    job = wait_for_idle(runner)
    assert job.state == "stopped"
    assert list(directory.glob("audit_*.xlsx"))     # partial, but written


def test_stop_with_no_job_is_harmless(ui):
    base, _runner, _dir = ui
    assert post(base, "/api/stop")["stopped"] is False


# --- reports listing --------------------------------------------------------


def test_reports_listing_is_empty_at_first(ui):
    base, _runner, _dir = ui
    assert get(base, "/api/reports")["reports"] == []


def test_report_appears_in_the_listing_after_a_crawl(ui, site):
    base, runner, _dir = ui
    post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 2,
                              "spellcheck": False})
    wait_for_idle(runner)

    reports = get(base, "/api/reports")["reports"]
    assert reports and reports[0]["name"].startswith("audit_")
    assert reports[0]["path"].endswith(".xlsx")


def test_open_rejects_a_missing_path(ui):
    base, _runner, _dir = ui
    result = post(base, "/api/open", {"path": "no-such-file.xlsx"})
    assert "error" in result


# --- keyword job ------------------------------------------------------------


def test_keyword_job_appends_sheets(ui, site):
    from openpyxl import load_workbook

    base, runner, directory = ui
    post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 2,
                              "spellcheck": False})
    wait_for_idle(runner)
    audit = next(iter(directory.glob("audit_*.xlsx")))

    csv_bytes = (Path(__file__).parent / "fixtures" / "gkp_english.csv").read_bytes()
    started = post(base, "/api/keywords?crawl=%s" % audit, raw=csv_bytes)
    assert started["started"] is True

    job = wait_for_idle(runner)
    assert job.state == "done", job.error
    assert "Gaps" in load_workbook(audit).sheetnames
    assert job.funnel and "rows in" in job.funnel
    assert set(job.stats) >= {"Keywords", "Gaps", "Owned"}


def test_keyword_job_without_a_crawl_is_reported(ui):
    base, runner, _dir = ui
    csv_bytes = (Path(__file__).parent / "fixtures" / "gkp_english.csv").read_bytes()
    post(base, "/api/keywords?crawl=nope.xlsx", raw=csv_bytes)
    job = wait_for_idle(runner)
    assert job.state == "error" and "audit spreadsheet" in job.error


# --- the one endpoint that touches the desktop ------------------------------


def test_open_is_restricted_to_this_machine(ui, monkeypatch):
    """/api/open hands a path to the OS, so a remote caller must be refused."""
    import seocrawl.ui as ui_module

    base, _runner, _dir = ui
    monkeypatch.setattr(ui_module.Handler, "_from_loopback", lambda self: False)
    result = post(base, "/api/open", {"path": "anything.xlsx"})
    assert "only allowed from this machine" in result["error"]


def test_open_is_allowed_from_loopback(ui, tmp_path, monkeypatch):
    import seocrawl.ui as ui_module

    base, _runner, _dir = ui
    opened = []

    def fake_open(path):
        opened.append(path)
        return {"opened": True, "revealed": False, "path": path}

    monkeypatch.setattr(ui_module, "open_in_os", fake_open)
    target = tmp_path / "report.xlsx"
    target.write_text("x", encoding="utf-8")
    assert post(base, "/api/open", {"path": str(target)})["opened"] is True
    assert opened == [str(target)]


def test_open_reveals_the_folder_when_no_app_is_registered(ui, tmp_path, monkeypatch):
    """os.startfile succeeds even with no handler, so "opened" would be a lie.

    A machine with no spreadsheet application installed is not unusual, and this
    one has none: .xlsx has no registered handler at all.
    """
    import seocrawl.ui as ui_module

    base, _runner, _dir = ui
    revealed = []
    monkeypatch.setattr(ui_module, "has_file_handler", lambda suffix: False)
    monkeypatch.setattr(ui_module, "reveal_in_file_manager", revealed.append)

    target = tmp_path / "report.xlsx"
    target.write_text("x", encoding="utf-8")
    result = post(base, "/api/open", {"path": str(target)})

    assert result["opened"] is False and result["revealed"] is True
    assert "registered to open" in result["note"]
    assert revealed == [target]


def test_open_uses_the_app_when_one_is_registered(ui, tmp_path, monkeypatch):
    import seocrawl.ui as ui_module

    base, _runner, _dir = ui
    monkeypatch.setattr(ui_module, "has_file_handler", lambda suffix: True)
    monkeypatch.setattr(ui_module.os, "startfile", lambda path: None, raising=False)
    monkeypatch.setattr(ui_module.subprocess, "Popen", lambda *a, **k: None)

    target = tmp_path / "report.html"
    target.write_text("<p>x</p>", encoding="utf-8")
    assert post(base, "/api/open", {"path": str(target)})["opened"] is True


# --- downloading reports ----------------------------------------------------


def test_report_downloads_as_an_attachment(ui, site):
    """The machine may have no spreadsheet app, so the browser gets the file."""
    base, runner, directory = ui
    post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 2,
                              "spellcheck": False})
    wait_for_idle(runner)

    xlsx = next(iter(directory.glob("audit_*.xlsx")))
    with urlopen(base + "/api/download?path=" + xlsx.name, timeout=15) as response:
        disposition = response.headers.get("Content-Disposition")
        payload = response.read()

    assert disposition == 'attachment; filename="%s"' % xlsx.name
    assert response.headers.get("Content-Type").endswith("spreadsheetml.sheet")
    assert payload == xlsx.read_bytes()


def test_html_report_is_viewable_inline(ui, site):
    base, runner, directory = ui
    post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 2,
                              "spellcheck": False})
    wait_for_idle(runner)

    html = next(iter(directory.glob("audit_*.html")))
    with urlopen(base + "/api/view?path=" + html.name, timeout=15) as response:
        assert response.headers.get("Content-Disposition") is None
        assert response.headers.get("Content-Type").startswith("text/html")
        assert b"SEO audit" in response.read()


def test_outputs_carry_browser_links(ui, site):
    base, runner, _dir = ui
    post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 2,
                              "spellcheck": False})
    job = wait_for_idle(runner)

    kinds = {out["kind"]: out for out in job.outputs}
    assert kinds["xlsx"]["download"].startswith("/api/download?path=")
    assert kinds["html"]["view"].startswith("/api/view?path=")


def test_reports_listing_carries_a_download_link(ui, site):
    base, runner, _dir = ui
    post(base, "/api/crawl", {"url": site, "delay": 0, "workers": 2,
                              "spellcheck": False})
    wait_for_idle(runner)
    assert get(base, "/api/reports")["reports"][0]["download"].startswith("/api/download")


def test_download_refuses_paths_outside_the_reports_folder(ui):
    """A plain GET is reachable from any page the browser visits."""
    base, _runner, _dir = ui
    for attempt in ("../../../../Windows/System32/drivers/etc/hosts",
                    "C:/Windows/win.ini", "/etc/passwd"):
        with pytest.raises(HTTPError) as excinfo:
            urlopen(base + "/api/download?path=" + attempt, timeout=10)
        assert excinfo.value.code == 404


def test_download_refuses_a_missing_file(ui):
    base, _runner, _dir = ui
    with pytest.raises(HTTPError) as excinfo:
        urlopen(base + "/api/download?path=nope.xlsx", timeout=10)
    assert excinfo.value.code == 404


def test_page_uses_download_links_not_open_buttons():
    assert "Download " in PAGE and "'View '" in PAGE
    assert "setAttribute('download'" in PAGE
