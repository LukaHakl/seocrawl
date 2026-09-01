"""End-to-end tests for ``seocrawl keywords``, including the written sheets."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from seocrawl.cli import main as seocrawl_main
from seocrawl.kwcli import main as keywords_main

FIXTURES = Path(__file__).parent / "fixtures"
ENGLISH = FIXTURES / "gkp_english.csv"
SLOVENIAN = FIXTURES / "gkp_slovenian.csv"
S = "https://shop.test"


@pytest.fixture(autouse=True)
def isolated_cwd(tmp_path, monkeypatch):
    """Run from an empty directory.

    ``Config.load`` falls back to ``./seocrawl.toml``, so without this the
    repo's own config would feed into tests meant to exercise the defaults.
    """
    workdir = tmp_path / "cwd"
    workdir.mkdir()
    monkeypatch.chdir(workdir)


@pytest.fixture
def audit(tmp_path) -> Path:
    """An audit workbook standing in for a real crawl."""
    from seocrawl.crawler import PageResult
    from seocrawl.export import write_workbook
    from seocrawl.findings import SiteStats

    pages = [
        PageResult(url=S + "/collections/bifold-wallets", status=200,
                   canonical_state="self", title="Bifold Wallets | Full Grain Leather",
                   h1="Bifold Wallets", inlinks=9),
        PageResult(url=S + "/products/traditional-bifold-leather-wallet", status=200,
                   canonical_state="self", title="Traditional Bifold Leather Wallet",
                   h1="Traditional Bifold Leather Wallet", product_group_id="TRAD"),
        PageResult(url=S + "/collections/card-holders", status=200,
                   canonical_state="self", title="Leather Card Holders", h1="Card Holders"),
        PageResult(url=S + "/blogs/guides/how-to-care-for-leather", status=200,
                   canonical_state="self", title="How to Care for Leather",
                   h1="How to Care for Leather"),
        PageResult(url=S + "/", status=200, canonical_state="self",
                   title="Handmade Leather Goods", h1="Handmade Leather Goods"),
    ]
    return write_workbook(tmp_path / "audit_shop-test.xlsx", pages, [], SiteStats())


def run(args, audit=None, **extra):
    """Invoke the subcommand with the common flags filled in."""
    argv = ["--gkp", str(extra.pop("gkp", ENGLISH))]
    if audit is not None:
        argv += ["--crawl", str(audit)]
    return keywords_main(argv + list(args))


# --- dispatch ---------------------------------------------------------------


def test_bare_url_form_still_works(capsys):
    """Adding a subcommand must not break "seocrawl <url>"."""
    assert seocrawl_main([]) == 2
    assert "usage: seocrawl" in capsys.readouterr().out


def test_keywords_subcommand_is_dispatched(audit, capsys):
    assert seocrawl_main(["keywords", "--gkp", str(ENGLISH), "--crawl", str(audit)]) == 0
    assert "Keyword map" in capsys.readouterr().out


def test_keywords_help_does_not_hit_the_crawl_parser(capsys):
    with pytest.raises(SystemExit) as excinfo:
        seocrawl_main(["keywords", "--help"])
    assert excinfo.value.code == 0
    assert "Keyword Planner" in capsys.readouterr().out


# --- running ----------------------------------------------------------------


def test_full_run_writes_three_sheets(audit, capsys):
    assert run([], audit) == 0
    book = load_workbook(audit)
    assert {"Keyword Map", "Gaps", "Cannibalization"} <= set(book.sheetnames)
    assert {"Summary", "Pages", "Findings"} <= set(book.sheetnames)   # audit intact


def test_console_prints_the_funnel_and_verdict(audit, capsys):
    run(["--contains", "wallet", "--brand", "popov"], audit)
    output = capsys.readouterr().out
    assert "rows in" in output and "unique" in output
    assert "OWNED" in output and "GAP" in output


def test_volume_caveat_is_always_printed(audit, capsys):
    run([], audit)
    output = capsys.readouterr().out
    assert "relative sizes, not truth" in output
    assert "GSC data supersedes" in output


def test_geo_note_is_included_in_the_caveat(audit, capsys):
    run(["--geo-note", "US"], audit)
    assert "Geo: US" in capsys.readouterr().out


def test_ingest_details_are_reported(audit, capsys):
    run([], audit, gkp=SLOVENIAN)
    output = capsys.readouterr().out
    assert "utf-16" in output and "tab-delimited" in output


def test_rerunning_does_not_stack_duplicate_sheets(audit):
    run([], audit)
    run([], audit)
    names = load_workbook(audit).sheetnames
    assert names.count("Keyword Map") == 1 and names.count("Gaps") == 1


def test_standalone_output_leaves_the_audit_alone(audit, tmp_path):
    output = tmp_path / "keyword_map.xlsx"
    assert run(["--standalone", "-o", str(output)], audit) == 0
    assert "Keyword Map" in load_workbook(output).sheetnames
    assert "Keyword Map" not in load_workbook(audit).sheetnames


# --- sheet contents ---------------------------------------------------------


def _sheet_rows(path, name):
    sheet = load_workbook(path)[name]
    header = [c.value for c in sheet[1]]
    return header, [dict(zip(header, row))
                    for row in sheet.iter_rows(min_row=2, values_only=True)]


def test_keyword_map_sheet_is_gap_first(audit):
    run([], audit)
    _header, rows = _sheet_rows(audit, "Keyword Map")
    order = {"GAP": 0, "CANNIBAL": 1, "WEAK": 2, "OWNED": 3}
    statuses = [order[r["Status"]] for r in rows]
    assert statuses == sorted(statuses)


def test_keyword_map_carries_volume_and_competition(audit):
    run([], audit)
    _header, rows = _sheet_rows(audit, "Keyword Map")
    bifold = next(r for r in rows if r["Keyword"] == "leather bifold wallet")
    assert bifold["Volume (mid)"] == 3162
    assert bifold["Volume range"] == "1K - 10K"
    assert bifold["Competition"] == "MED"
    assert bifold["Top bid (high)"] == 1.20


def test_cannibalization_sheet_lists_both_pages(audit):
    run([], audit)
    _header, rows = _sheet_rows(audit, "Cannibalization")
    contested = next(r for r in rows if r["Keyword"] == "leather bifold wallet")
    assert contested["Contenders"] == 2
    pages = {contested["Page A"], contested["Page B"]}
    assert pages == {S + "/collections/bifold-wallets",
                     S + "/products/traditional-bifold-leather-wallet"}
    assert "collection" in contested["Page types"] and "product" in contested["Page types"]


def test_gaps_sheet_has_a_suggested_action(audit):
    run([], audit)
    _header, rows = _sheet_rows(audit, "Gaps")
    assert rows
    for row in rows:
        assert "no owning URL" in row["Suggested action"]
        assert row["Keywords"] >= 1


def test_gap_keyword_appears_in_the_gaps_sheet(audit):
    run([], audit)
    _header, rows = _sheet_rows(audit, "Gaps")
    listed = " ".join(str(r["Keywords in cluster"]) for r in rows)
    assert "minimalist leather wallet" in listed


# --- filtering --------------------------------------------------------------


def test_brand_filter_removes_branded_keywords(audit):
    run(["--brand", "popov"], audit)
    _header, rows = _sheet_rows(audit, "Keyword Map")
    assert not any("popov" in r["Keyword"] for r in rows)


def test_contains_and_not_contains_combine(audit):
    run(["--contains", "wallet", "--not-contains", "code"], audit)
    _header, rows = _sheet_rows(audit, "Keyword Map")
    assert rows
    assert all("wallet" in r["Keyword"] for r in rows)
    assert not any("code" in r["Keyword"] for r in rows)


def test_min_volume_filter(audit):
    run(["--min-volume", "1000"], audit)
    _header, rows = _sheet_rows(audit, "Keyword Map")
    assert rows and all(r["Volume (mid)"] >= 1000 for r in rows)


def test_filtering_everything_out_is_an_error(audit, capsys):
    assert run(["--contains", "nothing-matches-this"], audit) == 1
    assert "No keywords left" in capsys.readouterr().out


# --- failure modes ----------------------------------------------------------


def test_missing_gkp_file_is_reported(audit, capsys, tmp_path):
    code = keywords_main(["--gkp", str(tmp_path / "nope.csv"), "--crawl", str(audit)])
    assert code == 2
    assert "Could not read the keyword file" in capsys.readouterr().out


def test_missing_crawl_file_is_reported(capsys, tmp_path):
    code = keywords_main(["--gkp", str(ENGLISH), "--crawl", str(tmp_path / "nope.xlsx")])
    assert code == 2
    assert "not found" in capsys.readouterr().out


def test_no_crawl_source_is_reported(capsys):
    assert keywords_main(["--gkp", str(ENGLISH)]) == 2
    assert "Pass --crawl" in capsys.readouterr().out


def test_unreadable_keyword_file_is_reported(capsys, tmp_path, audit):
    bad = tmp_path / "bad.csv"
    bad.write_text("Report\nApples,Oranges\n1,2\n", encoding="utf-8")
    assert keywords_main(["--gkp", str(bad), "--crawl", str(audit)]) == 2
    assert "No keyword column" in capsys.readouterr().out


# --- both exports agree end to end ------------------------------------------


def test_both_locales_produce_the_same_map(audit, tmp_path):
    english_out = tmp_path / "en.xlsx"
    slovenian_out = tmp_path / "sl.xlsx"
    run(["--standalone", "-o", str(english_out)], audit, gkp=ENGLISH)
    run(["--standalone", "-o", str(slovenian_out)], audit, gkp=SLOVENIAN)

    _h1, english = _sheet_rows(english_out, "Keyword Map")
    _h2, slovenian = _sheet_rows(slovenian_out, "Keyword Map")
    assert [(r["Keyword"], r["Status"], r["Volume (mid)"]) for r in english] == \
           [(r["Keyword"], r["Status"], r["Volume (mid)"]) for r in slovenian]


# --- the --no-idf escape hatch ----------------------------------------------


def test_no_idf_changes_the_verdict(audit, tmp_path, capsys):
    """Weighing every token equally lets a generic page claim a specific term."""
    with_idf = tmp_path / "idf.xlsx"
    without = tmp_path / "flat.xlsx"
    run(["--standalone", "-o", str(with_idf)], audit)
    run(["--standalone", "-o", str(without), "--no-idf"], audit)

    _h1, weighted = _sheet_rows(with_idf, "Keyword Map")
    _h2, flat = _sheet_rows(without, "Keyword Map")

    def status(rows, keyword):
        return next(r["Status"] for r in rows if r["Keyword"] == keyword)

    assert status(weighted, "minimalist leather wallet") == "GAP"
    assert status(flat, "minimalist leather wallet") == "WEAK"
